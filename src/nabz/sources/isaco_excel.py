"""خواندن فایل اکسل «لیست قیمت» ایساکو.

فایلی که از https://b2b.isaco.ir/PriceList دانلود می‌شود یک شیت به نام
«لیست قیمت» با هشت ستون دارد و حدود ۴۷ هزار ردیف:

    A کد تدارکاتی       ← ۱۰ رقمی، یکتا
    B شرح تدارکاتی
    C مصرف خودرو        ← ۴۳ مدل خودرو
    D واحد سنجش
    E تعداد در بسته
    F قیمت خرده فروشی    ← «قیمت مصرف کننده» در صفحهٔ وب
    G قیمت با ارزش افزوده ← F × ۱.۱
    H قیمت عمده          ← «في فروش» در صفحهٔ وب، F = H × ۱.۱۲

اگر ایساکو روزی ستون‌ها را عوض کند، :func:`load` با خطای روشن متوقف
می‌شود تا هرگز قیمت غلط روی سایت نرود.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterator

from openpyxl import load_workbook

from ..skucode import normalize

#: عنوان ستون‌هایی که حتماً باید وجود داشته باشند، به ترتیب.
EXPECTED_HEADERS: tuple[str, ...] = (
    "کد تدارکاتی",
    "شرح تدارکاتی",
    "مصرف خودرو",
    "واحد سنجش",
    "تعداد در بسته",
    "قیمت خرده فروشی",
    "قیمت با ارزش افزوده",
    "قیمت عمده",
)

#: اگر فایل کمتر از این تعداد ردیف داشت، احتمالاً ناقص دانلود شده.
MIN_EXPECTED_ROWS = 1_000


class IsacoFileError(RuntimeError):
    """فایل ایساکو ساختار مورد انتظار را ندارد."""


@dataclass(frozen=True, slots=True)
class IsacoItem:
    """یک ردیف از لیست قیمت ایساکو. همهٔ قیمت‌ها به ریال."""

    code: str          # کد نرمال‌شدهٔ ۱۰ رقمی
    name: str
    car: str           # مصرف خودرو
    unit: str          # واحد سنجش
    pack_size: int     # تعداد در بسته
    retail: int        # قیمت خرده‌فروشی (مصرف‌کننده)
    vat: int           # قیمت با ارزش افزوده
    wholesale: int     # قیمت عمده (فی فروش)

    @property
    def raw_code(self) -> str:
        """کد بدون صفرهای ابتدایی — همان چیزی که در حاسب ذخیره شده."""
        return self.code.lstrip("0") or "0"


def _as_int(value: object) -> int:
    """سلول عددی اکسل را به عدد صحیح تبدیل می‌کند."""
    if value is None or value == "":
        return 0
    if isinstance(value, (int, float)):
        return int(round(value))
    try:
        return int(round(float(str(value).replace(",", "").strip())))
    except ValueError:
        return 0


def _as_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _check_headers(row: tuple[object, ...]) -> None:
    """سطر عنوان را با ساختار مورد انتظار می‌سنجد."""
    actual = tuple(_as_text(c) for c in row[: len(EXPECTED_HEADERS)])
    if actual != EXPECTED_HEADERS:
        missing = [h for h in EXPECTED_HEADERS if h not in actual]
        raise IsacoFileError(
            "ساختار فایل ایساکو تغییر کرده است.\n"
            f"ستون‌های مورد انتظار: {', '.join(EXPECTED_HEADERS)}\n"
            f"ستون‌های موجود    : {', '.join(actual)}\n"
            + (f"ستون‌های گم‌شده    : {', '.join(missing)}" if missing else "")
        )


def iter_items(path: str | Path) -> Iterator[IsacoItem]:
    """ردیف‌های فایل را یکی‌یکی می‌خواند بدون اینکه همه را در حافظه نگه دارد."""
    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)

        try:
            header = next(rows)
        except StopIteration:
            raise IsacoFileError("فایل ایساکو خالی است.") from None
        _check_headers(header)

        for row in rows:
            if not row or row[0] is None:
                continue
            code = normalize(row[0])
            if not code:
                continue
            yield IsacoItem(
                code=code,
                name=_as_text(row[1]),
                car=_as_text(row[2]),
                unit=_as_text(row[3]),
                pack_size=_as_int(row[4]),
                retail=_as_int(row[5]),
                vat=_as_int(row[6]),
                wholesale=_as_int(row[7]),
            )
    finally:
        workbook.close()


def load(path: str | Path, *, strict: bool = True) -> dict[str, IsacoItem]:
    """کل لیست قیمت را می‌خواند و به دیکشنری «کد → کالا» تبدیل می‌کند.

    :param strict: اگر روشن باشد، فایل کوچک یا بدون قیمت رد می‌شود.
        این محافظ جلوی این را می‌گیرد که یک دانلود ناقص، قیمت‌های
        سایت را خراب کند.
    """
    items: dict[str, IsacoItem] = {}
    for item in iter_items(path):
        items[item.code] = item

    if strict:
        if len(items) < MIN_EXPECTED_ROWS:
            raise IsacoFileError(
                f"فایل ایساکو فقط {len(items):,} ردیف دارد و ناقص به نظر می‌رسد "
                f"(انتظار حداقل {MIN_EXPECTED_ROWS:,} ردیف). "
                "دانلود را دوباره انجام دهید."
            )
        priced = sum(1 for i in items.values() if i.vat > 0)
        if priced == 0:
            raise IsacoFileError("هیچ‌کدام از ردیف‌های فایل ایساکو قیمت ندارند.")

    return items


def build_prefix_index(
    items: dict[str, IsacoItem], length: int = 8
) -> dict[str, list[IsacoItem]]:
    """ایندکس «پیشوند ۸ رقمی → کالاهای هم‌خانواده».

    برای صفحهٔ «تطبیق دستی» لازم است: وقتی کد دقیق یک محصول سایت در
    لیست ایساکو نیست، نسخه‌های دیگر همان قطعه را پیشنهاد می‌دهیم.
    """
    index: dict[str, list[IsacoItem]] = {}
    for item in items.values():
        index.setdefault(item.code[:length], []).append(item)
    for family in index.values():
        family.sort(key=lambda i: i.code)
    return index
