"""خواندن اکسل/CSV خروجی نرم‌افزار حسابداری برای ایمپورت موجودی/قیمت.

خروجی حسابداری ساختار ثابتی ندارد، پس به‌جای هاردکد کردن ستون‌ها، کاربر
یک‌بار تعیین می‌کند کدام ستون «کد»، کدام «موجودی» و کدام «قیمت» است
(نگاشت ستون). این نگاشت ذخیره می‌شود تا دفعات بعد نیازی به تکرار نباشد.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from ..skucode import normalize


@dataclass(frozen=True, slots=True)
class ColumnMapping:
    """کدام ستون فایل به چه معناست (بر اساس اندیس صفر-پایه)."""

    code: int
    quantity: int | None = None
    price: int | None = None


@dataclass(frozen=True, slots=True)
class AccountingRow:
    code: str                      # کد نرمال‌شده
    raw_code: str
    quantity: Decimal | None = None
    price: Decimal | None = None


def read_headers(path: str | Path) -> list[str]:
    """سطر اول فایل را به‌عنوان عنوان ستون‌ها می‌خواند (برای نگاشت دستی)."""
    rows = _iter_raw_rows(path)
    try:
        return [str(c) if c is not None else "" for c in next(rows)]
    except StopIteration:
        return []


def read_rows(path: str | Path, mapping: ColumnMapping, *, has_header: bool = True):
    """ردیف‌های داده را بر اساس نگاشت ستون‌ها می‌خواند."""
    rows = _iter_raw_rows(path)
    if has_header:
        next(rows, None)

    result: list[AccountingRow] = []
    for raw in rows:
        if not raw or mapping.code >= len(raw):
            continue
        raw_code = raw[mapping.code]
        code = normalize(raw_code)
        if not code:
            continue
        result.append(
            AccountingRow(
                code=code,
                raw_code=str(raw_code).strip(),
                quantity=_num(raw, mapping.quantity),
                price=_num(raw, mapping.price),
            )
        )
    return result


def _num(row, index) -> Decimal | None:
    if index is None or index >= len(row) or row[index] is None:
        return None
    value = str(row[index]).replace(",", "").strip()
    value = value.translate(str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩",
                                          "01234567890123456789"))
    if not value:
        return None
    try:
        return Decimal(value)
    except (InvalidOperation, ValueError):
        return None


def _iter_raw_rows(path: str | Path):
    """ردیف‌های خام فایل را می‌دهد؛ هم xlsx و هم csv را پشتیبانی می‌کند."""
    path = Path(path)
    if path.suffix.lower() in (".csv", ".txt"):
        with path.open(encoding="utf-8-sig", newline="") as handle:
            yield from csv.reader(handle)
    else:
        workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
        try:
            for row in workbook.worksheets[0].iter_rows(values_only=True):
                yield list(row)
        finally:
            workbook.close()
